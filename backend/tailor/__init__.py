"""Tailor the master resume for a specific job, output a DOCX.

Bullet selection + skills emphasis are deterministic (rules over tags and JD
keywords). Only the 1-2 sentence summary goes through the LLM.
"""
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from db.models import Job, MasterResume, get_session
from llm import call_llm_fast

TAILOR_WORKERS = 4
MAX_BULLETS_PER_ROLE = 5
MIN_BULLETS_PER_ROLE = 2
MAX_EMPHASIZED_SKILLS = 8

OUTPUT_DIR = Path(__file__).parent.parent / "output" / "resumes"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

SUMMARY_PROMPT = """Rewrite this candidate's resume summary in 1-2 sentences to emphasize what's most relevant to the job. Do NOT add facts not in the original. Return only the rewritten sentences, no JSON, no preamble, no quotes.

Original summary:
{original}

Job: {title} at {company}
JD excerpt:
{jd_snippet}"""


def _score_bullet(bullet: dict, jd_lower: str) -> float:
    tag_score = sum(
        1
        for tag in bullet.get("tags", [])
        if tag.lower().replace("-", " ") in jd_lower
    )
    bullet_words = {
        w.lower().strip(".,():;") for w in bullet["text"].split() if len(w) > 4
    }
    jd_words = set(jd_lower.split())
    text_overlap = len(bullet_words & jd_words) / max(len(bullet_words), 1)
    metric_bonus = 0.5 if bullet.get("metrics") else 0.0
    scope_bonus = {"org": 0.3, "team": 0.15, "individual": 0.0}.get(
        bullet.get("scope", ""), 0.0
    )
    return tag_score + text_overlap + metric_bonus + scope_bonus


def _select_bullets(role_bullets: list[dict], jd_text: str) -> list[str]:
    jd_lower = jd_text.lower()
    real_bullets = [b for b in role_bullets if not b["text"].startswith("REPLACE")]
    if not real_bullets:
        return []
    scored = sorted(
        ((b, _score_bullet(b, jd_lower)) for b in real_bullets),
        key=lambda x: x[1],
        reverse=True,
    )
    keep = [b for b, s in scored if s > 0][:MAX_BULLETS_PER_ROLE]
    if len(keep) < MIN_BULLETS_PER_ROLE:
        keep = [b for b, _ in scored[:MIN_BULLETS_PER_ROLE]]
    return [b["text"] for b in keep]


def _emphasized_skills(master: dict, jd_lower: str) -> list[str]:
    all_skills: list[str] = []
    for items in master.get("skills", {}).values():
        all_skills.extend(items)
    matched = [s for s in all_skills if s.lower() in jd_lower]
    return matched[:MAX_EMPHASIZED_SKILLS]


def _write_summary(job: Job, original_summary: str) -> str:
    jd_snippet = (job.jd_text or "")[:1500]
    prompt = SUMMARY_PROMPT.format(
        original=original_summary,
        title=job.title,
        company=job.company,
        jd_snippet=jd_snippet,
    )
    last_err = None
    for attempt in range(3):
        try:
            text = call_llm_fast(prompt, want_json=False).strip()
            if text.startswith('"') and text.endswith('"'):
                text = text[1:-1].strip()
            return text
        except RuntimeError as e:
            last_err = e
            time.sleep(2 ** attempt)
    print(f"    summary LLM failed ({last_err}); using master summary verbatim")
    return original_summary


def tailor_resume(job: Job, master: dict) -> dict:
    jd_text = job.jd_text or ""
    jd_lower = jd_text.lower()

    experience = []
    for role in master["experience"]:
        bullets = _select_bullets(role["bullets"], jd_text)
        if not bullets:
            continue
        experience.append(
            {
                "company": role["company"],
                "role": role["role"],
                "dates": role["dates"],
                "bullets": bullets,
            }
        )

    summary = _write_summary(job, master.get("summary", ""))
    skills_emphasized = _emphasized_skills(master, jd_lower)

    return {
        "summary": summary,
        "experience": experience,
        "skills_emphasized": skills_emphasized,
        "changes_summary": (
            f"Selected {sum(len(r['bullets']) for r in experience)} bullets across "
            f"{len(experience)} roles by tag + keyword overlap. Emphasized skills: "
            f"{', '.join(skills_emphasized) or 'none matched'}."
        ),
    }


def _is_real(v) -> bool:
    return bool(v) and not str(v).startswith("REPLACE")


def _section_header(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "808080")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def _right_aligned_line(doc, left_text: str, right_text: str, *, bold_left: bool, space_before: int, space_after: int):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    left_run = p.add_run(left_text)
    left_run.bold = bold_left
    p.add_run(f"\t{right_text}")


def render_docx(tailored: dict, master: dict, output_path: Path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    name_run = p.add_run(master["personal"]["name"])
    name_run.bold = True
    name_run.font.size = Pt(18)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(0)
    p_info = master["personal"]
    contact_parts = [
        v for v in (
            p_info.get("location"),
            p_info.get("email"),
            p_info.get("phone"),
            p_info.get("linkedin"),
            p_info.get("github"),
            p_info.get("portfolio"),
        ) if _is_real(v)
    ]
    contact.add_run(" · ".join(contact_parts)).font.size = Pt(9.5)

    _section_header(doc, "Summary")
    summary_p = doc.add_paragraph(tailored["summary"])
    summary_p.paragraph_format.space_after = Pt(2)

    _section_header(doc, "Experience")
    for role in tailored["experience"]:
        _right_aligned_line(
            doc,
            f"{role['role']}, {role['company']}",
            role["dates"],
            bold_left=True,
            space_before=6,
            space_after=2,
        )
        for bullet in role["bullets"]:
            b = doc.add_paragraph(bullet, style="List Bullet")
            b.paragraph_format.space_after = Pt(2)

    skills_list = master.get("skills", {})
    if skills_list:
        _section_header(doc, "Skills")
        for cat, items in skills_list.items():
            if not items:
                continue
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(2)
            sp.add_run(f"{cat.replace('_', ' ').title()}: ").bold = True
            sp.add_run(", ".join(items))

    if master.get("education"):
        _section_header(doc, "Education")
        for edu in master["education"]:
            _right_aligned_line(
                doc,
                f"{edu['degree']}, {edu['school']}",
                edu["dates"],
                bold_left=True,
                space_before=4,
                space_after=2,
            )

    doc.save(output_path)


def _tailor_one(job: Job, master: dict) -> tuple[dict, Path]:
    tailored = tailor_resume(job, master)
    filename = f"{job.company.replace(' ', '_')}_{job.id[:8]}.docx"
    output_path = OUTPUT_DIR / filename
    render_docx(tailored, master, output_path)
    return tailored, output_path


def tailor_all_pending(min_score: float = 7.0):
    session = get_session()
    session.expire_on_commit = False
    try:
        master_row = session.query(MasterResume).first()
        if not master_row:
            print("ERROR: No master resume in DB. Seed it first.")
            return
        master = master_row.data

        pending = session.query(Job).filter(
            Job.score >= min_score,
            Job.resume_docx_path.is_(None),
        ).all()
        total = len(pending)
        print(f"Tailoring {total} resumes with {TAILOR_WORKERS} workers...")

        done = 0
        with ThreadPoolExecutor(max_workers=TAILOR_WORKERS) as pool:
            fut_to_job = {pool.submit(_tailor_one, j, master): j for j in pending}
            for fut in as_completed(fut_to_job):
                job = fut_to_job[fut]
                try:
                    tailored, output_path = fut.result()
                    job.tailored_bullets = tailored
                    job.resume_docx_path = str(output_path)
                    session.commit()
                    done += 1
                    print(f"  [{done}/{total}] {job.company} / {job.title}: {output_path.name}")
                except Exception as e:
                    print(f"  FAILED {job.company} / {job.title}: {e}")
                    session.rollback()
    finally:
        session.close()
